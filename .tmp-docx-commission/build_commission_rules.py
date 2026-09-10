from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\projetos\projeto-celular")
OUT = ROOT / "docs" / "regras-pendentes-comissoes.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.12
for style_name, size in [("Title", 22), ("Subtitle", 11), ("Heading 1", 16), ("Heading 2", 12)]:
    s = styles[style_name]
    s.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
    s.paragraph_format.space_after = Pt(6)
title_style_pr = styles["Title"]._element.get_or_add_pPr()
title_style_borders = title_style_pr.find(qn("w:pBdr"))
if title_style_borders is not None:
    title_style_pr.remove(title_style_borders)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color="D9D9D9"):
    tbl_pr = table._tbl.tblPr
    elem = tbl_pr.find(qn("w:tblBorders"))
    if elem is None:
        elem = OxmlElement("w:tblBorders")
        tbl_pr.append(elem)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), color)
        elem.append(node)


def set_cell_margin(cell, top=75, start=105, bottom=75, end=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Cm(width)
        shade(cell, "24364B")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
        set_cell_margin(cell)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, (value, width) in enumerate(zip(row, widths)):
            cells[i].width = Cm(width)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2:
                shade(cells[i], "F3F6F9")
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            r.font.size = Pt(8.3)
            set_cell_margin(cells[i])
    borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)


title = doc.add_paragraph(style="Title")
title.add_run("Regras Pendentes do Módulo de Comissões")
title_pr = title._p.get_or_add_pPr()
title_borders = title_pr.find(qn("w:pBdr"))
if title_borders is not None:
    title_pr.remove(title_borders)
sub = doc.add_paragraph(style="Subtitle")
sub.add_run("Documento para validação do cliente e fechamento das regras de negócio")
sub.paragraph_format.space_after = Pt(14)

p = doc.add_paragraph()
r = p.add_run("Conclusão principal. ")
r.bold = True
p.add_run(
    "A estrutura de histórico de comissões já permite congelar base, percentual e valor por venda. "
    "O que ainda falta é aprovar como a comissão será liberada em vendas com troca, quais deduções podem ser aplicadas "
    "e como tratar cancelamentos, estornos e pagamentos já realizados."
)

doc.add_heading("1 O que já está definido no sistema", level=1)
bullet("Cada item vendido gera um movimento histórico de comissão, sem depender do percentual futuro do vendedor.")
bullet("A base usa a margem positiva do item: valor vendido menos o custo congelado no momento da venda.")
bullet("O cálculo é feito por item e arredondado em centavos antes da soma final.")
bullet("Existem valores imediato e diferido, além dos estados pendente, parcialmente liberada, liberada e cancelada.")
bullet("As deduções são registradas separadamente e podem ser canceladas sem apagar o histórico.")

doc.add_heading("2 Regras que ainda precisam da aprovação do cliente", level=1)
add_table(
    ["Tema", "Comportamento atual", "Regra recomendada", "Decisão necessária"],
    [
        ("Venda normal", "Comissão totalmente imediata.", "Manter 100 por cento liberada após a conclusão válida da venda.", "Confirmar o momento exato da liberação."),
        ("Venda com troca", "Ainda não existe divisão automática.", "Gerar a comissão completa, liberar uma parte e manter outra vinculada ao item recebido na troca.", "Definir a fórmula da parte imediata e da parte diferida."),
        ("Giro da troca", "Não há evento automático de liberação.", "Liberar a parte diferida quando o aparelho recebido for revendido em uma venda válida.", "Definir se basta revender ou se também deve estar pago e fora do prazo de devolução."),
        ("Base da comissão", "Margem do item sem custo negativo.", "Usar venda do item menos CMV congelado; taxas, impostos e despesas ficam fora, salvo decisão expressa.", "Confirmar se taxas de cartão também reduzem a base."),
        ("Estorno", "Precisa de regra completa.", "Cancelar apenas a comissão dos itens estornados e preservar os demais.", "Definir recuperação de comissão já paga."),
        ("Serviços e garantias", "Podem ter custo zero e gerar comissão integral.", "Criar regra de elegibilidade por categoria de item.", "Confirmar quais categorias comissionam."),
        ("Saldo negativo", "Sem regra fechada.", "Transportar o saldo para o próximo período, com histórico e motivo.", "Confirmar se haverá limite ou desconto em folha."),
    ],
    [3.0, 4.0, 6.1, 3.6],
)

doc.add_heading("3 Proposta para vendas com troca", level=1)
p = doc.add_paragraph()
p.add_run("Modelo recomendado. ").bold = True
p.add_run(
    "A comissão deve ser calculada normalmente por item, mas a parcela associada ao valor aceito na troca fica diferida. "
    "Ela só é liberada quando o produto recebido na troca volta a girar no estoque por meio de uma nova venda válida."
)
bullet("Parte imediata: comissão sobre a parcela da margem que não depende do valor recebido na troca.")
bullet("Parte diferida: comissão correspondente à parcela vinculada ao bem recebido como pagamento.")
bullet("Liberação: ocorre após a revenda válida do item de troca.")
bullet("Rastreabilidade: a comissão diferida deve guardar o item recebido e a venda que causou sua liberação.")

p = doc.add_paragraph()
p.add_run("Ponto que precisa ser escolhido. ").bold = True
p.add_run(
    "Se a empresa preferir uma regra mais simples, poderá usar uma divisão fixa, como um percentual imediato e outro diferido. "
    "O percentual não deve ser programado antes da aprovação do cliente."
)

doc.add_heading("4 Deduções da comissão", level=1)
p = doc.add_paragraph()
p.add_run("Regra geral. ").bold = True
p.add_run("Comissão líquida a pagar é a comissão liberada menos as deduções ativas do mesmo vendedor e período.")
add_table(
    ["Tipo de dedução", "Uso esperado", "Tratamento recomendado"],
    [
        ("Adiantamento ou vale", "Valor antecipado ao vendedor.", "Abater no próximo fechamento até liquidar o saldo."),
        ("Estorno", "Venda ou item cancelado após gerar comissão.", "Reverter o movimento; se já pago, lançar dedução no período seguinte."),
        ("Falta ou desconto", "Desconto autorizado pela empresa.", "Exigir motivo, responsável e data."),
        ("Ajuste manual", "Correção excepcional de cálculo.", "Exigir justificativa e manter trilha de auditoria."),
        ("Outros", "Caso não coberto pelos tipos anteriores.", "Usar apenas com descrição obrigatória."),
    ],
    [4.0, 5.4, 7.3],
)
p = doc.add_paragraph()
p.add_run("Importante. ").bold = True
p.add_run(
    "Dedução de comissão não é despesa operacional da empresa. Ela reduz o valor líquido devido ao vendedor e deve permanecer fora da DRE de despesas. "
    "Quando cancelada, a dedução continua no histórico, mas deixa de reduzir o pagamento."
)

doc.add_heading("5 Cancelamentos e situações especiais", level=1)
bullet("Estorno total da venda: cancelar todas as comissões ligadas à venda.")
bullet("Estorno parcial: cancelar somente os movimentos dos itens estornados.")
bullet("Comissão já paga: lançar uma dedução de estorno no próximo fechamento, sem apagar o pagamento anterior.")
bullet("Troca rejeitada ou devolvida: manter a parte diferida bloqueada até decisão administrativa.")
bullet("Produto de troca vendido com prejuízo: definir se a comissão diferida será reduzida, cancelada ou liberada integralmente.")
bullet("Produto de troca sem giro por longo período: definir prazo e quem pode autorizar liberação excepcional.")
bullet("Substituição de vendedor: manter o vendedor histórico da venda; mudanças futuras não alteram períodos fechados.")

doc.add_heading("6 Estados e datas do movimento", level=1)
add_table(
    ["Estado", "Significado", "Quando usar"],
    [
        ("Pendente", "Nenhum valor liberado.", "Comissão aguardando validação ou condição de negócio."),
        ("Parcialmente liberada", "Existe valor imediato e saldo diferido.", "Venda com troca ou outra condição futura."),
        ("Liberada", "Todo o valor pode entrar no fechamento.", "Venda normal ou condição diferida cumprida."),
        ("Cancelada", "O movimento não é mais devido.", "Estorno, cancelamento ou decisão autorizada."),
    ],
    [4.0, 5.6, 7.1],
)
bullet("Data de geração: quando a comissão é criada a partir da venda.")
bullet("Data de liberação: quando o valor se torna elegível ao fechamento.")
bullet("Data de pagamento: quando o vendedor efetivamente recebe; recomenda-se registrar separadamente.")

doc.add_heading("7 Informações adicionais recomendadas no banco", level=1)
p = doc.add_paragraph(
    "Para automatizar as regras sem perder rastreabilidade, a próxima evolução deverá relacionar cada comissão diferida ao item recebido na troca e à venda que liberou o valor. "
    "Também é recomendável registrar fechamento, pagamento, usuário responsável e motivo de cancelamento."
)
add_table(
    ["Informação", "Finalidade"],
    [
        ("Item recebido na troca", "Identificar exatamente qual produto mantém a comissão diferida."),
        ("Venda de liberação", "Comprovar qual revenda causou o giro e liberou o saldo."),
        ("Fechamento de comissão", "Agrupar valores por vendedor e competência."),
        ("Data e estado do pagamento", "Separar valor liberado de valor efetivamente pago."),
        ("Motivo e responsável", "Auditar ajustes, cancelamentos e liberações manuais."),
    ],
    [6.0, 10.7],
)

doc.add_heading("8 Checklist para aprovação", level=1)
for item in [
    "Definir a fórmula da parte imediata e da parte diferida em vendas com troca.",
    "Definir o evento que caracteriza o giro da troca.",
    "Confirmar se taxa de cartão, impostos ou outros custos reduzem a base.",
    "Confirmar quais categorias de produtos e serviços geram comissão.",
    "Definir o tratamento do item recebido e vendido com prejuízo.",
    "Definir como recuperar comissão já paga em uma venda estornada.",
    "Definir se saldo negativo passa para o período seguinte e quais limites se aplicam.",
    "Definir periodicidade e data de fechamento e pagamento.",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_after = Pt(5)
    p.add_run("☐  " + item)

doc.add_heading("9 Critérios de aceite", level=1)
bullet("Alterar o percentual de um vendedor não modifica comissões históricas.")
bullet("Alterar o custo do produto não modifica a base histórica da comissão.")
bullet("Uma venda normal gera comissão liberada conforme a regra aprovada.")
bullet("Uma venda com troca preserva e rastreia os valores imediato e diferido.")
bullet("O giro do produto de troca libera somente a comissão vinculada a ele.")
bullet("Estornos e deduções não apagam dados e deixam histórico de usuário, data e motivo.")
bullet("Os relatórios distinguem comissão gerada, liberada, paga, deduzida e líquida.")

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Próximo passo. ").bold = True
p.add_run("O cliente deve preencher o checklist acima. Com essas respostas, a regra pode ser implementada sem suposições que alterem o valor devido aos vendedores.")

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Start Seminovos  |  Regras de comissões  |  10 de setembro de 2026")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

doc.core_properties.title = "Regras Pendentes do Módulo de Comissões"
doc.core_properties.subject = "Regras de negócio para validação"
doc.core_properties.author = "Start Seminovos"
doc.save(OUT)
print(OUT)
